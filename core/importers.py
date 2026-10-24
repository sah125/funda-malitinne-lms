"""Turn face-to-face training material (Word .docx) into an online course.

A QCTO learner guide is structured with Word headings:
    Heading 1  -> topic           e.g. "KM-01-KT02: Project management framework"
    Heading 2  -> topic element   e.g. "KT0201 Definitions and key concepts"
Each *topic element* becomes one online lesson (a short, focused read), the
topic introduction is placed at the start of its first lesson, and all lessons
are grouped under one Knowledge Module. Tables, bullet lists, bold sub-headings
and pictures are kept.

The matching *learner assessment* (.docx) turns every "Activity N: IACxxxx"
into an online assignment that learners answer and upload — their formative
evidence for the Portfolio of Evidence.

Used by:  python manage.py import_learner_guide …   and the instructor
"Import learner guide" page.
"""
import html
import re
from dataclasses import dataclass, field
from datetime import timedelta

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

LESSON_MARKER = '<!--html-->'
WORDS_PER_MINUTE = 180


@dataclass
class Section:
    title: str
    parts: list = field(default_factory=list)   # html fragments
    words: int = 0
    topic: str = ''

    def html(self):
        return LESSON_MARKER + '\n' + '\n'.join(self.parts)


def _iter_block_items(document):
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            yield Paragraph(child, document)
        elif tag == 'tbl':
            yield Table(child, document)


def _is_list(paragraph):
    ppr = paragraph._p.pPr
    return (ppr is not None and ppr.numPr is not None) or paragraph.style.name.lower().startswith('list')


def _is_bold_subheading(paragraph):
    text = paragraph.text.strip()
    if not text or len(text) > 90 or text.endswith(('.', ':', ';', ',')):
        return False
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(r.bold or (r.style is not None and r.style.font.bold) for r in runs)


def _run_html(paragraph):
    out = []
    for run in paragraph.runs:
        text = html.escape(run.text)
        if not text:
            continue
        if run.bold:
            text = f'<strong>{text}</strong>'
        if run.italic:
            text = f'<em>{text}</em>'
        out.append(text)
    return ''.join(out).strip()


def _images(paragraph, document, folder):
    """Save pictures embedded in the paragraph and return <img> tags."""
    tags = []
    for blip in paragraph._p.xpath('.//a:blip'):
        rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        part = document.part.related_parts.get(rid) if rid else None
        if part is None or not getattr(part, 'blob', None):
            continue
        ext = (part.partname.ext or 'png').lower()
        if ext not in ('png', 'jpg', 'jpeg', 'gif', 'webp'):
            continue
        name = default_storage.save(f'{folder}/{slugify(rid)}.{ext}', ContentFile(part.blob))
        tags.append(f'<figure><img src="{default_storage.url(name)}" alt="Illustration" loading="lazy"></figure>')
    return tags


def _table_html(table):
    rows = []
    for i, row in enumerate(table.rows):
        cells = []
        seen = set()
        for cell in row.cells:
            if id(cell._tc) in seen:   # merged cells repeat
                continue
            seen.add(id(cell._tc))
            text = '<br>'.join(html.escape(p.text.strip()) for p in cell.paragraphs if p.text.strip())
            tag = 'th' if i == 0 else 'td'
            cells.append(f'<{tag}>{text}</{tag}>')
        rows.append('<tr>' + ''.join(cells) + '</tr>')
    if not rows:
        return ''
    return '<div class="table-wrap"><table>' + rows[0].join(['<thead>', '</thead>']) + '<tbody>' + ''.join(rows[1:]) + '</tbody></table></div>'


def parse_learner_guide(path_or_file, image_folder='lesson_images/import'):
    """Return a list of Section objects (one per online lesson)."""
    import docx
    document = docx.Document(path_or_file)
    sections, current, topic, topic_intro = [], None, '', []
    list_open = False

    def close_list():
        nonlocal list_open
        if list_open and current is not None:
            current.parts.append('</ul>')
        list_open = False

    started = False
    for block in _iter_block_items(document):
        if block.__class__.__name__ == 'Table':
            if not started:
                continue
            target = current.parts if current else topic_intro
            close_list()
            target.append(_table_html(block))
            continue

        para = block
        style = para.style.name
        text = para.text.strip()

        if style == 'Heading 1' and text:
            started = True
            close_list()
            topic = re.sub(r'\s*\(\d+%\)\s*$', '', text)
            current = Section(title=topic, topic=topic)
            sections.append(current)
            topic_intro = current.parts
            continue
        if not started:
            continue
        if style == 'Heading 2' and text:
            close_list()
            parent = sections[-1] if sections else None
            current_words = parent.words if parent else 0
            # Topic heading with no own text yet -> its intro moves into this first lesson
            if parent and parent.title == topic and parent is current:
                sections.pop()
                intro = [f'<h3>{html.escape(topic)}</h3>'] + parent.parts
            else:
                intro = []
            current = Section(title=re.sub(r'\s+', ' ', text), topic=topic, words=current_words if intro else 0)
            if intro:
                current.parts.append('<div class="lesson-topic-intro">' + '\n'.join(intro) + '</div>')
            sections.append(current)
            continue
        if current is None:
            continue

        imgs = _images(para, document, image_folder)
        if imgs:
            close_list()
            current.parts.extend(imgs)
        if not text:
            continue
        current.words += len(text.split())
        if style.startswith('Heading'):
            close_list()
            current.parts.append(f'<h3>{html.escape(text)}</h3>')
        elif _is_bold_subheading(para):
            close_list()
            current.parts.append(f'<h3>{html.escape(text)}</h3>')
        elif _is_list(para):
            if not list_open:
                current.parts.append('<ul>')
                list_open = True
            current.parts.append(f'<li>{_run_html(para) or html.escape(text)}</li>')
        else:
            close_list()
            current.parts.append(f'<p>{_run_html(para) or html.escape(text)}</p>')
    close_list()
    return [s for s in sections if s.parts]


@dataclass
class Activity:
    title: str
    topic: str
    question: str


def parse_learner_assessment(path_or_file):
    import docx
    document = docx.Document(path_or_file)
    activities, topic, current = [], '', None
    for para in document.paragraphs:
        text = para.text.strip()
        style = para.style.name
        if style == 'Heading 1' and text:
            topic = text
            current = None
        elif style == 'Heading 2' and text.lower().startswith('activity'):
            current = Activity(title=re.sub(r'\s+', ' ', text), topic=topic, question='')
            activities.append(current)
        elif current and text and not re.fullmatch(r'[_\s.]+', text):
            current.question += (('\n' if current.question else '') + text)
    return [a for a in activities if a.question]


@transaction.atomic
def import_into_course(course, guide_file, assessment_file=None, module_title='', replace=False, due_in_days=30):
    """Create lessons (and assignments) in `course`. Returns a summary dict."""
    from .models import Assignment, LearningModule, Lesson

    folder = f'lesson_images/course_{course.pk}'
    sections = parse_learner_guide(guide_file, image_folder=folder)
    if not sections:
        raise ValueError('No headings found. Make sure the learner guide uses Word "Heading 1" for topics and "Heading 2" for topic elements.')

    module, _ = LearningModule.objects.get_or_create(
        course=course, title=module_title or 'Knowledge Module',
        defaults={'module_type': 'knowledge', 'order': course.learning_modules.count() + 1,
                  'description': 'Imported from the face-to-face learner guide.'},
    )
    if replace:
        Lesson.objects.filter(course=course, module=module).delete()

    start = (course.lessons.order_by('-order').values_list('order', flat=True).first() or 0) + 1
    created = 0
    for offset, section in enumerate(sections):
        minutes = max(5, round(section.words / WORDS_PER_MINUTE))
        Lesson.objects.create(
            course=course, module=module, title=section.title[:200], content=section.html(),
            duration=minutes, order=start + offset,
        )
        created += 1

    assignments = 0
    if assessment_file:
        due = timezone.now() + timedelta(days=due_in_days)
        for act in parse_learner_assessment(assessment_file):
            topic = re.sub(r'\s*\(\d+%\)\s*$', '', act.topic)
            Assignment.objects.get_or_create(
                course=course, title=f'{act.title.strip()} — {topic}'[:200],
                defaults={'description': f'{act.question}\n\nAnswer in your own words. Type your answer in a Word document '
                                         f'or PDF and upload it here. It forms part of your Portfolio of Evidence.',
                          'due_date': due, 'total_points': 100},
            )
            assignments += 1

    return {'lessons': created, 'assignments': assignments, 'module': module.title,
            'minutes': sum(max(5, round(s.words / WORDS_PER_MINUTE)) for s in sections)}
